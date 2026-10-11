from __future__ import annotations

import csv
import json
import subprocess
import tempfile
import unittest
import zipfile
from pathlib import Path

import numpy as np
import rasterio

try:  # openpyxl is optional (not a declared dependency); the xlsx audit table needs it
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


ROOT = Path(__file__).resolve().parents[1]
PYTHON = ROOT / ".venv/bin/python"
SCRIPT = ROOT / "bin/visualize/plot_secocha_standardization_detail.py"
BASE = ROOT / "outputs/secocha_pim_chain"


@unittest.skipIf(load_workbook is None, "openpyxl is not installed; the script writes an xlsx table")
class PeruSecochaStandardizationDetailTest(unittest.TestCase):
    def test_builds_source_backed_standardization_detail(self):
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary)
            completed = subprocess.run(
                [str(PYTHON), str(SCRIPT), "--output-dir", str(output)],
                cwd=ROOT,
                capture_output=True,
                text=True,
                check=False,
            )
            self.assertEqual(completed.returncode, 0, completed.stderr)

            process_png = output / "secocha_standardization_process.png"
            process_svg = output / "secocha_standardization_process.svg"
            xlsx_path = output / "IFPM_Standardization_Audit.xlsx"
            csv_path = output / "standardization_detail.csv"
            manifest_path = output / "standardization_detail_manifest.json"
            self.assertEqual(xlsx_path.name, "IFPM_Standardization_Audit.xlsx")
            self.assertGreater(process_png.stat().st_size, 100_000)
            self.assertGreater(process_svg.stat().st_size, 20_000)
            self.assertGreater(xlsx_path.stat().st_size, 5_000)

            with csv_path.open(newline="", encoding="utf-8") as handle:
                rows = {row["Indicator code"]: row for row in csv.DictReader(handle)}
            self.assertEqual(len(rows), 14)
            self.assertEqual(
                rows["twi_wetness"]["Standardization method"],
                "Percentile min–max scaling (2nd–98th percentiles)",
            )
            self.assertEqual(
                rows["drainage_proximity"]["Direction"],
                "Lower measurement → higher score",
            )
            self.assertEqual(
                rows["road_access_proximity"]["Standardization method"],
                "Distance-decay scaling",
            )
            self.assertEqual(float(rows["road_access_proximity"]["Upper bound"]), 20_000.0)
            self.assertEqual(
                rows["land_cover_pressure"]["Standardization method"],
                "Categorical lookup-table scaling",
            )
            self.assertEqual(
                rows["built_agriculture_road_expansion"]["Standardization method"],
                "Weighted combination",
            )
            self.assertEqual(
                rows["protected_ecological_context"]["Standardization method"],
                "Verified-absence rule",
            )
            self.assertTrue(all(row["Original unit"].strip() for row in rows.values()))
            self.assertEqual(
                rows["rainfall_climatology"]["Original unit"],
                "millimetres per December–March season",
            )
            self.assertEqual(rows["twi_wetness"]["Original unit"], "Dimensionless")
            self.assertEqual(
                rows["soil_k_factor"]["Original unit"],
                "t·ha·h/(ha·MJ·mm)",
            )
            self.assertEqual(rows["ls_factor"]["Original unit"], "Dimensionless")
            self.assertEqual(rows["c_factor"]["Original unit"], "Dimensionless")
            self.assertEqual(
                rows["settlement_population"]["Original unit"],
                "People per 30-m grid cell",
            )
            self.assertEqual(
                rows["settlement_population"]["Source dataset"],
                "WorldPop population raster (2020)",
            )
            self.assertNotIn("<!-- Raw unit -->", process_svg.read_text(encoding="utf-8"))
            process_svg_text = process_svg.read_text(encoding="utf-8")
            self.assertNotIn("A7367", process_svg_text)
            self.assertNotIn("transfer curves", process_svg_text)
            with zipfile.ZipFile(xlsx_path) as workbook:
                self.assertIn("xl/worksheets/sheet1.xml", workbook.namelist())
                workbook_text = "\n".join(
                    workbook.read(name).decode("utf-8", errors="replace")
                    for name in workbook.namelist()
                    if name.endswith(".xml")
                )
                worksheet_xml = [
                    workbook.read(name).decode("utf-8")
                    for name in workbook.namelist()
                    if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
                ]
                table_xml = [
                    workbook.read(name).decode("utf-8")
                    for name in workbook.namelist()
                    if name.startswith("xl/tables/table") and name.endswith(".xml")
                ]
            self.assertTrue(all("<autoFilter" not in text for text in worksheet_xml))
            self.assertEqual(sum("<autoFilter" in text for text in table_xml), 5)
            self.assertIn("Original unit", workbook_text)
            self.assertIn("Source dataset", workbook_text)
            self.assertIn("soil_k_factor", workbook_text)
            self.assertIn("Study area", workbook_text)
            self.assertIn("region-independent", workbook_text.lower())
            self.assertNotIn("Standardization template", workbook_text)
            self.assertNotIn("supervisor_full_spec", workbook_text)
            self.assertNotIn("secocha_reference_extent", workbook_text)
            self.assertNotIn("outputs/", workbook_text)
            self.assertNotIn("data/raw", workbook_text)
            self.assertNotIn("Claim level", workbook_text)
            self.assertNotIn("LUT", workbook_text)
            self.assertNotIn("lookup_table", workbook_text)
            self.assertNotIn("D_max", workbook_text)
            self.assertNotIn("raw units", workbook_text.lower())
            self.assertNotIn("class_specific_rule", workbook_text)
            self.assertIn("Topographic Wetness Index (TWI)", workbook_text)
            self.assertIn("Slope–length factor (LS)", workbook_text)
            self.assertIn("Cover-management factor (C)", workbook_text)
            workbook = load_workbook(xlsx_path)
            sheet = workbook["Standardization detail"]
            self.assertEqual(sheet.freeze_panes, "A5")
            self.assertEqual(sheet.max_row, 18)
            self.assertEqual(sheet.max_column, 13)
            self.assertEqual(sheet.page_setup.orientation, "landscape")
            self.assertEqual(sheet.page_setup.fitToWidth, 1)
            self.assertEqual(sheet.print_title_rows, "$1:$4")
            self.assertLessEqual(sheet.column_dimensions["M"].width, 24)
            self.assertIn("StandardizationDetail", sheet.tables)
            self.assertFalse(any("secocha" in name.lower() for name in sheet.tables))
            detail_table = sheet.tables["StandardizationDetail"]
            self.assertEqual(detail_table.tableStyleInfo.name, "TableStyleLight1")
            self.assertFalse(detail_table.tableStyleInfo.showRowStripes)
            self.assertTrue(sheet["A4"].fill.fgColor.rgb.endswith("D9EAF7"))
            self.assertTrue(
                all(
                    cell.fill.fill_type is None
                    for row in sheet.iter_rows(min_row=5, max_row=18, min_col=1, max_col=13)
                    for cell in row
                )
            )
            detail_text = "\n".join(
                str(cell.value)
                for row in sheet.iter_rows()
                for cell in row
                if cell.value is not None
            )
            self.assertNotIn("clip(", detail_text)
            self.assertNotIn("Source dataset", detail_text)

            self.assertIn("Read me", workbook.sheetnames)
            self.assertIn("Explanations", workbook.sheetnames)
            self.assertIn("Data sources", workbook.sheetnames)
            self.assertIn("Additional materials", workbook.sheetnames)
            self.assertIn("Rule catalog", workbook.sheetnames)
            self.assertNotIn("Standardization template", workbook.sheetnames)
            for sheet_name in ("Read me", "Standardization detail", "Explanations", "Data sources", "Additional materials", "Rule catalog"):
                self.assertTrue(workbook[sheet_name].sheet_view.showGridLines)

            explanations = workbook["Explanations"]
            self.assertEqual(explanations.max_row, 18)
            self.assertEqual(explanations.max_column, 6)
            self.assertIn("StandardizationExplanations", explanations.tables)
            self.assertFalse(explanations.tables["StandardizationExplanations"].tableStyleInfo.showRowStripes)
            self.assertTrue(
                all(
                    cell.fill.fill_type is None
                    for row in explanations.iter_rows(min_row=5, max_row=18, min_col=1, max_col=6)
                    for cell in row
                )
            )
            explanation_text = "\n".join(
                str(cell.value)
                for row in explanations.iter_rows()
                for cell in row
                if cell.value is not None
            )
            self.assertIn("Formula", explanation_text)
            self.assertIn("Rule description", explanation_text)
            self.assertIn("Source dataset", explanation_text)
            self.assertIn("clip(", explanation_text)

            sources = workbook["Data sources"]
            self.assertEqual(sources.max_row, 12)
            self.assertEqual(sources.max_column, 7)
            self.assertIn("StandardizationDataSources", sources.tables)
            self.assertFalse(sources.tables["StandardizationDataSources"].tableStyleInfo.showRowStripes)
            source_text = "\n".join(
                str(cell.value)
                for row in sources.iter_rows()
                for cell in row
                if cell.value is not None
            )
            self.assertIn("Source dataset", source_text)
            self.assertIn("Used for indicators", source_text)
            self.assertIn("Source portal", source_text)
            self.assertIn("Download URL", source_text)
            self.assertIn("WorldPop Peru 2020 population count", source_text)
            self.assertIn("Geofabrik OpenStreetMap Peru extract", source_text)
            self.assertIn("https://data.worldpop.org/GIS/Population/Global_2000_2020/2020/PER/per_ppp_2020_UNadj.tif", source_text)
            self.assertIn("https://download.geofabrik.de/south-america/peru-latest-free.shp.zip", source_text)
            self.assertIn("https://data.chc.ucsb.edu/products/CHIRPS/v3.0/", source_text)
            self.assertTrue(sources["D5"].hyperlink)
            self.assertTrue(sources["E5"].hyperlink)
            self.assertTrue(
                all(
                    all(part.startswith("https://") for part in str(row[4].value).splitlines())
                    for row in sources.iter_rows(min_row=5, max_row=12)
                )
            )
            self.assertNotIn("/Users/", source_text)
            self.assertTrue(
                all(
                    cell.fill.fill_type is None
                    for row in sources.iter_rows(min_row=5, max_row=12, min_col=1, max_col=7)
                    for cell in row
                )
            )

            additional = workbook["Additional materials"]
            self.assertEqual(additional.max_row, 14)
            self.assertEqual(additional.max_column, 7)
            self.assertIn("StandardizationAdditionalMaterials", additional.tables)
            self.assertFalse(additional.tables["StandardizationAdditionalMaterials"].tableStyleInfo.showRowStripes)
            additional_text = "\n".join(
                str(cell.value)
                for row in additional.iter_rows()
                for cell in row
                if cell.value is not None
            )
            self.assertIn("Active study-area boundary", additional_text)
            self.assertIn("INGEMMET small-scale mining overlay", additional_text)
            self.assertIn("excluded from the priority score", additional_text)
            self.assertIn("Criteria weights and combination rules", additional_text)
            self.assertIn("Standardized indicator rasters", additional_text)
            self.assertIn("INGEMMET A7367 hazard assessment report", additional_text)
            self.assertIn("INDECI Emergency Report 1733", additional_text)
            self.assertIn("CHIRPS v3 event-window rainfall", additional_text)
            self.assertIn("Copernicus Sentinel before/after catalogue", additional_text)
            self.assertIn("INDECI SINPAD emergency records", additional_text)
            self.assertIn("https://sigrid.cenepred.gob.pe/sigridv3/storage/biblioteca/", additional_text)
            self.assertIn("https://portal.indeci.gob.pe/wp-content/uploads/2023/05/", additional_text)
            self.assertIn("https://data.chc.ucsb.edu/products/CHIRPS/v3.0/daily/final/sat/2023/", additional_text)
            self.assertIn("https://stac.dataspace.copernicus.eu/v1/collections/sentinel-1-grd/items/", additional_text)
            self.assertIn("https://geosinpad.indeci.gob.pe/indeci/rest/services/Emergencias/", additional_text)
            self.assertIn("https://geocatminnube.ingemmet.gob.pe/arcgis/rest/services/SERV_PEQUENA_MINERIA/MapServer/0", additional_text)
            self.assertTrue(additional["E5"].hyperlink)
            self.assertTrue(additional["E6"].hyperlink)
            self.assertTrue(additional["E10"].hyperlink)
            self.assertTrue(additional["E11"].hyperlink)
            self.assertTrue(additional["E12"].hyperlink)
            self.assertTrue(additional["E13"].hyperlink)
            self.assertTrue(additional["E14"].hyperlink)
            self.assertNotIn("/Users/", additional_text)
            self.assertTrue(
                all(
                    cell.fill.fill_type is None
                    for row in additional.iter_rows(min_row=5, max_row=14, min_col=1, max_col=7)
                    for cell in row
                )
            )

            rules = workbook["Rule catalog"]
            rule_text = "\n".join(
                str(cell.value)
                for row in rules.iter_rows()
                for cell in row
                if cell.value is not None
            )
            for method in (
                "Percentile min–max scaling",
                "Distance-decay scaling",
                "Categorical lookup-table scaling",
                "Binary class rule",
                "Weighted combination",
                "Verified-absence rule",
            ):
                self.assertIn(method, rule_text)
            self.assertIn("StandardizationRuleCatalog", rules.tables)
            self.assertNotIn("Secocha", rule_text)

            readme_text = "\n".join(
                str(cell.value)
                for row in workbook["Read me"].iter_rows()
                for cell in row
                if cell.value is not None
            )
            self.assertIn("Standardization detail", readme_text)
            self.assertIn("Explanations", readme_text)
            self.assertIn("Data sources", readme_text)
            self.assertIn("Evidence status", readme_text)
            self.assertNotIn("Claim level", readme_text)
            readme = workbook["Read me"]
            self.assertTrue(readme.sheet_view.showGridLines)
            self.assertEqual(readme.page_setup.orientation, "landscape")
            self.assertEqual(readme.page_setup.fitToWidth, 1)
            self.assertGreaterEqual(readme.column_dimensions["A"].width, 34)
            for row_index in range(3, 15):
                self.assertEqual(readme.cell(row_index, 1).alignment.horizontal, "left")
                self.assertEqual(readme.cell(row_index, 1).alignment.vertical, "center")
                self.assertEqual(readme.cell(row_index, 2).alignment.horizontal, "left")
                self.assertEqual(readme.cell(row_index, 2).alignment.vertical, "center")

            twi_path = BASE / (
                "hydrology_derivatives/"
                "secocha_twi.tif"
            )
            with rasterio.open(twi_path) as source:
                twi = source.read(1).astype("float64")
                valid = np.isfinite(twi) & (twi != source.nodata)
            expected_low, expected_high = np.percentile(twi[valid], [2, 98])
            self.assertAlmostEqual(float(rows["twi_wetness"]["Lower bound"]), expected_low, places=6)
            self.assertAlmostEqual(float(rows["twi_wetness"]["Upper bound"]), expected_high, places=6)

            for row in rows.values():
                self.assertGreaterEqual(float(row["Standardized minimum"]), 0.0)
                self.assertLessEqual(float(row["Standardized maximum"]), 1.0)
                self.assertGreater(float(row["Valid-cell fraction"]), 0.0)

            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            self.assertTrue(manifest["public_safe"])
            self.assertEqual(manifest["status"], "Preliminary example")
            self.assertNotIn("boundary_id", manifest)
            self.assertNotIn("experiment_id", manifest)
            self.assertNotIn("claim_level", manifest)
            self.assertEqual(manifest["indicator_count"], 14)
            self.assertTrue(manifest["separated_outputs"])
            self.assertTrue(manifest["generic_table_artifact"])
            self.assertTrue(all("_" not in key for key in manifest["method_counts"]))
            self.assertNotIn("/", json.dumps(manifest))
            self.assertEqual(manifest["table_artifact_name"], "IFPM_Standardization_Audit.xlsx")
            self.assertEqual(
                manifest["figures"]["process"]["png"],
                "secocha_standardization_process.png",
            )
            self.assertEqual(
                manifest["figures"]["table"]["xlsx"],
                "IFPM_Standardization_Audit.xlsx",
            )
            self.assertEqual(manifest["table_format"], "xlsx")
            self.assertEqual(manifest["example_worksheet"], "Standardization detail")
            self.assertEqual(manifest["rule_catalog_worksheet"], "Rule catalog")
            self.assertEqual(manifest["explanation_worksheet"], "Explanations")
            self.assertEqual(manifest["explanation_row_count"], 14)
            self.assertEqual(manifest["source_worksheet"], "Data sources")
            self.assertEqual(manifest["source_row_count"], 8)
            self.assertEqual(manifest["source_download_url_column"], "Download URL")
            self.assertEqual(manifest["additional_worksheet"], "Additional materials")
            self.assertEqual(manifest["additional_row_count"], 10)
            self.assertGreaterEqual(manifest["rule_catalog_count"], 6)
            self.assertEqual(manifest["layout_card_overflow_count"], 0)
            self.assertEqual(
                manifest["visible_table_columns"],
                [
                    "Indicator code", "Indicator group", "Indicator", "Original unit",
                    "Standardization method", "Direction", "Lower bound", "Upper bound",
                    "Standardized minimum", "Standardized maximum", "Standardized mean",
                    "Valid-cell fraction", "Evidence status",
                ],
            )
            self.assertEqual(manifest["layout_overlap_count"], 0)
            self.assertGreaterEqual(manifest["layout_checked_text_count"], 8)
            self.assertGreaterEqual(manifest["layout_min_vertical_gap_px"], 8.0)
            self.assertGreaterEqual(manifest["layout_title_subtitle_gap_px"], 8.0)
            self.assertGreaterEqual(manifest["layout_footer_gap_px"], 8.0)
            self.assertIn("different indicators require different standardization methods", manifest["takeaway"])


if __name__ == "__main__":
    unittest.main()
