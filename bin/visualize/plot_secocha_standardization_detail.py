#!/usr/bin/env python3
"""Render a public-facing standardization detail workbook and source-backed example.

This is a local OUTPUT generator. It reads the active Secocha rasters and
manifests, writes a terminology-safe indicator table, and exports a report-
friendly PNG/SVG companion. It does not alter model scores.
"""

from __future__ import annotations

import argparse
import csv
import json
from dataclasses import dataclass
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import rasterio
from matplotlib.patches import FancyBboxPatch

try:  # optional dependency, only needed for the xlsx audit table
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError:
    Workbook = None


ROOT = Path(__file__).resolve().parents[2]
BOUNDARY_ID = "secocha_reference_extent"
OUTPUT_ROOT_ID = "secocha_pim_chain"  # outputs folder name (differs from the id since 2026-10-11 rename)
EXPERIMENT_ID = "supervisor_full_spec_v07_secocha"
STUDY_AREA = "Secocha"
BASE = ROOT / f"outputs/{OUTPUT_ROOT_ID}"
CRITERIA = BASE / "criteria_stack/criteria"
HYDRO = BASE / "hydrology_derivatives"
PRIORITY = BASE / "priority_scenarios"
DEFAULT_OUTPUT = PRIORITY / "figures"
EXPLANATION_COLUMNS = [
    ("Indicator code", "component_id"),
    ("Indicator group", "criterion_group"),
    ("Indicator", "display_name"),
    ("Original unit", "raw_unit"),
    ("Standardization method", "method"),
    ("Direction", "direction"),
    ("Lower bound", "lower_bound"),
    ("Upper bound", "upper_bound"),
    ("Formula", "formula"),
    ("Rule description", "rule_detail"),
    ("Standardized minimum", "score_min"),
    ("Standardized maximum", "score_max"),
    ("Standardized mean", "score_mean"),
    ("Valid-cell fraction", "valid_fraction"),
    ("Evidence status", "claim_level"),
    ("Source dataset", "raw_source_path"),
    ("Standardized output", "score_path"),
]
SUMMARY_COLUMNS = [
    ("Indicator code", "component_id"),
    ("Indicator group", "criterion_group"),
    ("Indicator", "display_name"),
    ("Original unit", "raw_unit"),
    ("Standardization method", "method"),
    ("Direction", "direction"),
    ("Lower bound", "lower_bound"),
    ("Upper bound", "upper_bound"),
    ("Standardized minimum", "score_min"),
    ("Standardized maximum", "score_max"),
    ("Standardized mean", "score_mean"),
    ("Valid-cell fraction", "valid_fraction"),
    ("Evidence status", "claim_level"),
]
EXPLANATION_SHEET_COLUMNS = [
    ("Indicator code", "component_id"),
    ("Indicator", "display_name"),
    ("Formula", "formula"),
    ("Rule description", "rule_detail"),
    ("Source dataset", "raw_source_path"),
    ("Standardized output", "score_path"),
]
SOURCE_COLUMNS = [
    ("Source ID", "source_id"),
    ("Source dataset", "source_dataset"),
    ("Used for indicators", "used_for"),
    ("Source portal", "source_portal"),
    ("Download URL", "download_url"),
    ("Access / processing note", "source_note"),
    ("Evidence status", "claim_level"),
]
ADDITIONAL_COLUMNS = [
    ("Material ID", "material_id"),
    ("Material", "material"),
    ("Material type", "material_type"),
    ("Role in model", "model_role"),
    ("Source or download URL", "source_url"),
    ("Status / use", "status_use"),
    ("Note", "note"),
]
ADDITIONAL_ROWS = [
    {
        "material_id": "MAT-01",
        "material": "Active study-area boundary",
        "material_type": "Spatial mask and reference extent",
        "model_role": "Defines common support and clipping for all indicators",
        "source_url": (
            "https://sigrid.cenepred.gob.pe/sigridv3/storage/biblioteca/"
            "16130_informe-tecnico-n0-a7367-evaluacion-de-peligros-geologicos-en-en-"
            "centro-poblado-secocha-y-las-localidades-la-poderosa-san-martin-posco-"
            "misky-y-pampayl.pdf"
        ),
        "status_use": "Used for spatial masking; no standalone public boundary file",
        "note": "Reference extent interpreted from the public A7367 report map; not a hydrological catchment or event footprint.",
    },
    {
        "material_id": "MAT-02",
        "material": "INGEMMET small-scale mining overlay",
        "material_type": "External contextual overlay",
        "model_role": "Provides separate mining context outside the priority score",
        "source_url": "https://geocatminnube.ingemmet.gob.pe/arcgis/rest/services/SERV_PEQUENA_MINERIA/MapServer/0",
        "status_use": "Context only; excluded from the priority score",
        "note": "Service query recorded 26 features on 2026-08-09; the count may change and this overlay is not a scored indicator.",
    },
    {
        "material_id": "MAT-03",
        "material": "Criteria weights and combination rules",
        "material_type": "Model parameter",
        "model_role": "Combines the five standardized criteria into a priority score",
        "source_url": "Method specification; no external download",
        "status_use": "Model parameter; not a source dataset",
        "note": "Weights and missing-data rules define how standardized layers are combined.",
    },
    {
        "material_id": "MAT-04",
        "material": "Standardized indicator rasters",
        "material_type": "Derived intermediate output",
        "model_role": "Provides the 14 comparable 0–1 layers used by the criteria",
        "source_url": "Listed source datasets; no external download",
        "status_use": "Intermediate model output",
        "note": "These rasters are derived products, not independent source datasets.",
    },
    {
        "material_id": "MAT-05",
        "material": "Criterion and priority rasters",
        "material_type": "Derived model output",
        "model_role": "Produces the five criterion layers and the prototype priority result",
        "source_url": "Derived model output; no external download",
        "status_use": "Prototype output; not a final flood-risk claim",
        "note": "The outputs depend on the standardized indicators, declared weights, and common-support policy.",
    },
    {
        "material_id": "MAT-06",
        "material": "INGEMMET A7367 hazard assessment report",
        "material_type": "Historical hazard reference",
        "model_role": "Provides mapped geological-hazard context and affected localities; does not enter the priority score",
        "source_url": (
            "https://sigrid.cenepred.gob.pe/sigridv3/storage/biblioteca/"
            "16130_informe-tecnico-n0-a7367-evaluacion-de-peligros-geologicos-en-en-"
            "centro-poblado-secocha-y-las-localidades-la-poderosa-san-martin-posco-"
            "misky-y-pampayl.pdf"
        ),
        "status_use": "Public report; validation context only",
        "note": "Official report used to interpret the reference extent and historical hazard context.",
    },
    {
        "material_id": "MAT-07",
        "material": "INDECI Emergency Report 1733",
        "material_type": "Historical event record",
        "model_role": "Provides event chronology and reported impacts; does not enter the priority score",
        "source_url": (
            "https://portal.indeci.gob.pe/wp-content/uploads/2023/05/"
            "INFORME-DE-EMERGENCIA-N%C2%BA-1733-20MAY2023-LLUVIAS-INTENSAS-EN-EL-"
            "DISTRITO-DE-MARIANO-NICOLAS-VALC%C3%81RCEL-AREQUIPA-13-DEE.pdf"
        ),
        "status_use": "Public report; validation context only",
        "note": "Report dated 20 May 2023 for Mariano Nicolás Valcárcel and nearby localities.",
    },
    {
        "material_id": "MAT-08",
        "material": "CHIRPS v3 event-window rainfall",
        "material_type": "Event rainfall context",
        "model_role": "Provides daily rainfall context around the historical event; separate from the long-term criterion",
        "source_url": (
            "https://data.chc.ucsb.edu/products/CHIRPS/v3.0/daily/final/sat/2023/"
            "chirps-v3.0.sat.2023.02.05.tif"
        ),
        "status_use": "Public download; validation context only",
        "note": "Daily files cover 2023-01-07 to 2023-02-10; native 0.05-degree grid, not 30-m evidence. Product portal: https://data.chc.ucsb.edu/products/CHIRPS/v3.0/",
    },
    {
        "material_id": "MAT-09",
        "material": "Copernicus Sentinel before/after catalogue",
        "material_type": "Satellite validation context",
        "model_role": "Provides scene metadata and previews for before/after comparison; does not enter the priority score",
        "source_url": (
            "https://stac.dataspace.copernicus.eu/v1/collections/sentinel-1-grd/items/"
            "S1A_IW_GRDH_1SDV_20230129T102426_20230129T102458_046997_05A329_0328_COG"
        ),
        "status_use": "Metadata and previews public; pixel download requires an account",
        "note": "Selected scenes are catalogue context; quantitative Sentinel pixels were not downloaded. Catalogue search: https://stac.dataspace.copernicus.eu/v1/search; product download requires Copernicus Data Space access.",
    },
    {
        "material_id": "MAT-10",
        "material": "INDECI SINPAD emergency records",
        "material_type": "Emergency-event reference data",
        "model_role": "Provides event-point validation context; does not enter the priority score",
        "source_url": (
            "https://geosinpad.indeci.gob.pe/indeci/rest/services/Emergencias/"
            "EMERGENCIAS_SINPAD/MapServer/0"
        ),
        "status_use": "Public ArcGIS service; query results may change",
        "note": "Flood, rainfall, and debris-flow records support validation checks, not proof of a final flood-risk result.",
    },
]
RULE_CATALOG_COLUMNS = [
    "Method",
    "Standardization method",
    "Use when",
    "Parameters to record",
    "Formula or mapping",
    "Interpretation",
    "Original-unit requirement",
    "Missing-data handling",
    "Required evidence in output",
]
RULE_CATALOG_ROWS = [
    [
        "Percentile min–max scaling",
        "Scale continuous measurements using the 2nd and 98th percentiles",
        "Continuous variables with local outliers or skew",
        "2nd percentile, 98th percentile, and direction",
        "clip((x - p_2) / (p_98 - p_2), 0, 1)",
        "Higher or lower measurement → higher score",
        "State the original measurement unit",
        "Missing and outside-study-area cells remain NoData",
        "0–1 score, bounds, and percentile evidence",
    ],
    [
        "Distance-decay scaling",
        "Scale distance to a feature using a declared influence range",
        "Distance-to-feature indicators with a declared influence range",
        "Maximum influence distance (D)",
        "clip(1 - distance / D, 0, 1)",
        "Shorter distance → higher score",
        "State the distance unit",
        "Missing and outside-study-area cells remain NoData",
        "0–1 score, D, and feature source",
    ],
    [
        "Categorical lookup-table scaling",
        "Assign scores to categorical classes using an explicit table",
        "Categorical classes with an explicit risk or pressure value",
        "Complete class-to-score lookup table",
        "score = lookup table(class)",
        "Class-based interpretation",
        "State the class coding system",
        "Unmapped classes remain NoData unless explicitly assigned",
        "0–1 score and the complete lookup table",
    ],
    [
        "Binary class rule",
        "Assign 1 to a selected class or condition and 0 to the declared fallback",
        "Presence/absence or one selected class drives the indicator",
        "Selected class and fallback rule",
        "score = 1[class = selected_class]",
        "Class-based interpretation",
        "State the class coding system",
        "Unmapped or missing source cells remain NoData",
        "0/1 score, selected class, and source",
    ],
    [
        "Weighted combination",
        "Combine already standardized indicators into one indicator group",
        "Several already standardized indicators represent one indicator group",
        "Indicator codes and weights; weights sum to 1",
        "score = sum(weight_i × score_i)",
        "Higher indicator scores → higher combined score",
        "Inputs must already be comparable 0–1 scores",
        "Declare missing-data propagation; never silently fill with zero",
        "0–1 score, indicator weights, and coverage",
    ],
    [
        "Verified-absence rule",
        "Assign zero only when a defined source confirms no mapped feature",
        "A queried source explicitly confirms no mapped feature",
        "Source, query extent, date, and absence evidence",
        "score = 0 only for a verified absence in the queried source",
        "Verified absence → 0; unknown → NoData",
        "State the feature dataset and query scope",
        "Unknown or unqueried areas remain NoData",
        "0 score and source-backed absence evidence",
    ],
]

PUBLIC_METHOD_LABELS = {
    "p2_p98_minmax": "Percentile min–max scaling (2nd–98th percentiles)",
    "fixed_distance_decay": "Distance-decay scaling",
    "class_value_lookup": "Categorical lookup-table scaling",
    "binary_class_lookup": "Binary class rule",
    "weighted_composite": "Weighted combination",
    "confirmed_zero_context": "Verified-absence rule",
}
PUBLIC_DIRECTION_LABELS = {
    "higher_raw_higher_score": "Higher measurement → higher score",
    "lower_raw_higher_score": "Lower measurement → higher score",
    "class_specific_rule": "Class-based interpretation",
    "higher_component_higher_score": "Higher indicator scores → higher combined score",
    "confirmed_absence_in_queried_layer": "Verified absence → 0; unknown → NoData",
}
PUBLIC_RAW_UNITS = {
    "rainfall_climatology": "millimetres per December–March season",
    "twi_wetness": "Dimensionless",
    "drainage_proximity": "metres",
    "soil_k_factor": "t·ha·h/(ha·MJ·mm)",
    "ls_factor": "Dimensionless",
    "c_factor": "Dimensionless",
    "land_cover_pressure": "ESA WorldCover class code (categorical)",
    "bare_soil_vegetation_loss": "ESA WorldCover class code (categorical)",
    "built_agriculture_road_expansion": "Dimensionless indicator scores (0–1)",
    "water_wetland_context": "Metres from ESA WorldCover water/wetland classes",
    "protected_ecological_context": "Protected-area feature presence (binary)",
    "road_access_proximity": "Metres from the OpenStreetMap road network",
    "settlement_population": "People per 30-m grid cell",
    "critical_infrastructure": "ESA WorldCover class code (categorical)",
}
PUBLIC_SOURCE_LABELS = {
    "rainfall_climatology": "Seasonal rainfall climatology (December–March, 2014–2024)",
    "twi_wetness": "Topographic Wetness Index (TWI) raster",
    "drainage_proximity": "Distance-to-drainage raster",
    "soil_k_factor": "Soil erodibility factor (K) raster",
    "ls_factor": "Slope–length factor (LS) raster",
    "c_factor": "Cover-management factor (C) raster",
    "land_cover_pressure": "ESA WorldCover land-cover raster",
    "bare_soil_vegetation_loss": "ESA WorldCover land-cover raster",
    "built_agriculture_road_expansion": "Built-up, agriculture, and road indicator layers",
    "water_wetland_context": "ESA WorldCover water and wetland classes",
    "protected_ecological_context": "National protected-area dataset",
    "road_access_proximity": "OpenStreetMap road network",
    "settlement_population": "WorldPop population raster (2020)",
    "critical_infrastructure": "ESA WorldCover built-up class",
}
PUBLIC_SOURCE_ROLES = {
    "rainfall_climatology": "Raw rainfall measurement",
    "twi_wetness": "Derived terrain indicator",
    "drainage_proximity": "Derived distance indicator",
    "soil_k_factor": "Soil property measurement",
    "ls_factor": "Derived terrain factor",
    "c_factor": "Land-cover factor",
    "land_cover_pressure": "Land-cover class map",
    "bare_soil_vegetation_loss": "Land-cover class map",
    "built_agriculture_road_expansion": "Combined indicator layers",
    "water_wetland_context": "Land-cover classes used for distance derivation",
    "protected_ecological_context": "Feature-presence dataset",
    "road_access_proximity": "Road network used for distance derivation",
    "settlement_population": "Population measurement",
    "critical_infrastructure": "Land-cover class map",
}
PUBLIC_SOURCE_NOTES = {
    "rainfall_climatology": "Seasonal December–March summary for 2014–2024; used as the rainfall measurement before standardization.",
    "twi_wetness": "Terrain-derived wetness indicator on the working reference grid.",
    "drainage_proximity": "Distance to mapped drainage on the working reference grid.",
    "soil_k_factor": "Soil erodibility factor aligned to the working reference grid; accessed from the official SoilGrids collection.",
    "ls_factor": "Slope–length factor on the working reference grid.",
    "c_factor": "Cover-management factor derived from land-cover information on the working reference grid.",
    "land_cover_pressure": "Land-cover classes are mapped to explicit runoff-pressure values.",
    "bare_soil_vegetation_loss": "Land-cover classes are used to identify bare or sparse vegetation.",
    "built_agriculture_road_expansion": "Combines built-up, agriculture, and road indicator layers.",
    "water_wetland_context": "Water and wetland classes are used to derive a proximity indicator.",
    "protected_ecological_context": "Queried national protected-area source; absence is interpreted only for this example.",
    "road_access_proximity": "Road features are used to derive distance to the road network.",
    "settlement_population": "Population values from 2020 are used before percentile scaling.",
    "critical_infrastructure": "Built-up land-cover class is used as the infrastructure indicator.",
}
PUBLIC_SOURCE_PORTALS = {
    "rainfall_climatology": "https://www.chc.ucsb.edu/data/chirps3\nhttps://data.chc.ucsb.edu/products/CHIRPS/v3.0/",
    "twi_wetness": "https://doi.org/10.5066/F7PR7TFT",
    "drainage_proximity": "https://doi.org/10.5066/F7PR7TFT",
    "soil_k_factor": "https://www.isric.org/explore/soilgrids",
    "ls_factor": "https://doi.org/10.5066/F7PR7TFT",
    "c_factor": "https://esa-worldcover.org/en/data-access",
    "land_cover_pressure": "https://esa-worldcover.org/en/data-access",
    "bare_soil_vegetation_loss": "https://esa-worldcover.org/en/data-access",
    "built_agriculture_road_expansion": "https://esa-worldcover.org/en/data-access\nhttps://download.geofabrik.de/south-america/peru.html",
    "water_wetland_context": "https://esa-worldcover.org/en/data-access",
    "protected_ecological_context": "https://geoservicios.sernanp.gob.pe/arcgis/rest/services/servicios_ogc/peru_sernanp_0102/MapServer/0",
    "road_access_proximity": "https://download.geofabrik.de/south-america/peru.html",
    "settlement_population": "https://hub.worldpop.org/",
    "critical_infrastructure": "https://esa-worldcover.org/en/data-access",
}
PUBLIC_DOWNLOAD_URLS = {
    "rainfall_climatology": "https://data.chc.ucsb.edu/products/CHIRPS/v3.0/",
    "twi_wetness": "https://step.esa.int/auxdata/dem/SRTMGL1/S16W074.SRTMGL1.hgt.zip",
    "drainage_proximity": "https://step.esa.int/auxdata/dem/SRTMGL1/S16W074.SRTMGL1.hgt.zip",
    "soil_k_factor": "https://files.isric.org/soilgrids/latest/data/clay/clay_0-5cm_mean.vrt\nhttps://files.isric.org/soilgrids/latest/data/silt/silt_0-5cm_mean.vrt\nhttps://files.isric.org/soilgrids/latest/data/sand/sand_0-5cm_mean.vrt\nhttps://files.isric.org/soilgrids/latest/data/soc/soc_0-5cm_mean.vrt",
    "ls_factor": "https://step.esa.int/auxdata/dem/SRTMGL1/S16W074.SRTMGL1.hgt.zip",
    "c_factor": "https://esa-worldcover.s3.eu-central-1.amazonaws.com/v200/2021/map/ESA_WorldCover_10m_2021_v200_S18W075_Map.tif",
    "land_cover_pressure": "https://esa-worldcover.s3.eu-central-1.amazonaws.com/v200/2021/map/ESA_WorldCover_10m_2021_v200_S18W075_Map.tif",
    "bare_soil_vegetation_loss": "https://esa-worldcover.s3.eu-central-1.amazonaws.com/v200/2021/map/ESA_WorldCover_10m_2021_v200_S18W075_Map.tif",
    "built_agriculture_road_expansion": "https://esa-worldcover.s3.eu-central-1.amazonaws.com/v200/2021/map/ESA_WorldCover_10m_2021_v200_S18W075_Map.tif\nhttps://download.geofabrik.de/south-america/peru-latest-free.shp.zip",
    "water_wetland_context": "https://esa-worldcover.s3.eu-central-1.amazonaws.com/v200/2021/map/ESA_WorldCover_10m_2021_v200_S18W075_Map.tif",
    "protected_ecological_context": "https://geoservicios.sernanp.gob.pe/arcgis/rest/services/servicios_ogc/peru_sernanp_0102/MapServer/0",
    "road_access_proximity": "https://download.geofabrik.de/south-america/peru-latest-free.shp.zip",
    "settlement_population": "https://data.worldpop.org/GIS/Population/Global_2000_2020/2020/PER/per_ppp_2020_UNadj.tif",
    "critical_infrastructure": "https://esa-worldcover.s3.eu-central-1.amazonaws.com/v200/2021/map/ESA_WorldCover_10m_2021_v200_S18W075_Map.tif",
}
PUBLIC_SOURCE_KEYS = {
    "rainfall_climatology": "chirps_v3",
    "twi_wetness": "srtm_dem",
    "drainage_proximity": "srtm_dem",
    "soil_k_factor": "soilgrids",
    "ls_factor": "srtm_dem",
    "c_factor": "worldcover",
    "land_cover_pressure": "worldcover",
    "bare_soil_vegetation_loss": "worldcover",
    "built_agriculture_road_expansion": "worldcover_osm",
    "water_wetland_context": "worldcover",
    "protected_ecological_context": "sernanp",
    "road_access_proximity": "geofabrik_osm",
    "settlement_population": "worldpop",
    "critical_infrastructure": "worldcover",
}
PUBLIC_SOURCE_REGISTRY = {
    "chirps_v3": {
        "source_id": "SRC-01",
        "source_dataset": "CHIRPS v3 rainfall",
        "source_portal": PUBLIC_SOURCE_PORTALS["rainfall_climatology"],
        "download_url": PUBLIC_DOWNLOAD_URLS["rainfall_climatology"],
        "source_note": "Monthly values were summarized for December–March (2014–2024) and resampled to the working grid; no new spatial information is implied.",
    },
    "srtm_dem": {
        "source_id": "SRC-02",
        "source_dataset": "SRTMGL1 v3 elevation",
        "source_portal": PUBLIC_SOURCE_PORTALS["twi_wetness"],
        "download_url": PUBLIC_DOWNLOAD_URLS["twi_wetness"],
        "source_note": "The elevation tile is the download source for the terrain wetness, drainage-distance, and slope–length derivatives.",
    },
    "soilgrids": {
        "source_id": "SRC-03",
        "source_dataset": "SoilGrids 250 m soil properties",
        "source_portal": PUBLIC_SOURCE_PORTALS["soil_k_factor"],
        "download_url": PUBLIC_DOWNLOAD_URLS["soil_k_factor"],
        "source_note": "Official SoilGrids data were accessed through the remote collection for the soil-erodibility calculation.",
    },
    "worldcover": {
        "source_id": "SRC-04",
        "source_dataset": "ESA WorldCover 10 m 2021 v200",
        "source_portal": PUBLIC_SOURCE_PORTALS["c_factor"],
        "download_url": PUBLIC_DOWNLOAD_URLS["c_factor"],
        "source_note": "The 2021 land-cover classes support the cover factor, class rules, and water/wetland proximity indicator.",
    },
    "worldcover_osm": {
        "source_id": "SRC-05",
        "source_dataset": "ESA WorldCover plus Geofabrik OpenStreetMap",
        "source_portal": PUBLIC_SOURCE_PORTALS["built_agriculture_road_expansion"],
        "download_url": PUBLIC_DOWNLOAD_URLS["built_agriculture_road_expansion"],
        "source_note": "Land-cover and road indicators are combined for the built-up, agriculture, and road pressure example.",
    },
    "geofabrik_osm": {
        "source_id": "SRC-06",
        "source_dataset": "Geofabrik OpenStreetMap Peru extract",
        "source_portal": PUBLIC_SOURCE_PORTALS["road_access_proximity"],
        "download_url": PUBLIC_DOWNLOAD_URLS["road_access_proximity"],
        "source_note": "The Peru road extract is used to derive distance to the road network.",
    },
    "sernanp": {
        "source_id": "SRC-07",
        "source_dataset": "SERNANP national protected areas, MapServer layer 0",
        "source_portal": PUBLIC_SOURCE_PORTALS["protected_ecological_context"],
        "download_url": PUBLIC_DOWNLOAD_URLS["protected_ecological_context"],
        "source_note": "The queried layer returned no features within the example query scope; this does not cover unqueried protected-area categories.",
    },
    "worldpop": {
        "source_id": "SRC-08",
        "source_dataset": "WorldPop Peru 2020 population count",
        "source_portal": PUBLIC_SOURCE_PORTALS["settlement_population"],
        "download_url": PUBLIC_DOWNLOAD_URLS["settlement_population"],
        "source_note": "The 2020 UN-adjusted population raster is used before aggregation and percentile scaling.",
    },
}
PUBLIC_OUTPUT_LABEL = "Standardized indicator raster (0–1)"
PUBLIC_STATUS = "Preliminary example"


@dataclass(frozen=True)
class ComponentSpec:
    component_id: str
    display_name: str
    group: str
    method: str
    direction: str
    raw_unit: str
    score_path: Path
    rule: str
    raw_path: Path | None = None
    lower_bound: float | None = None
    upper_bound: float | None = None
    formula: str = ""


def read_array(path: Path) -> np.ndarray:
    with rasterio.open(path) as source:
        values = source.read(1).astype("float64")
        if source.nodata is not None:
            values[values == source.nodata] = np.nan
    values[~np.isfinite(values)] = np.nan
    return values


def percentile_bounds(path: Path) -> tuple[float, float]:
    values = read_array(path)
    valid = values[np.isfinite(values)]
    if not valid.size:
        raise ValueError(f"no valid cells for percentile bounds: {path}")
    low, high = np.percentile(valid, [2, 98])
    if not np.isfinite(low) or not np.isfinite(high) or high <= low:
        raise ValueError(f"invalid p2-p98 bounds for {path}: {low}, {high}")
    return float(low), float(high)


def continuous_spec(
    component_id: str,
    display_name: str,
    group: str,
    raw_path: Path,
    score_path: Path,
    raw_unit: str,
    *,
    invert: bool = False,
    bounds: tuple[float, float] | None = None,
) -> ComponentSpec:
    low, high = bounds or percentile_bounds(raw_path)
    direction = "lower_raw_higher_score" if invert else "higher_raw_higher_score"
    formula = (
        "1 - clip((x - p2) / (p98 - p2), 0, 1)"
        if invert
        else "clip((x - p2) / (p98 - p2), 0, 1)"
    )
    return ComponentSpec(
        component_id=component_id,
        display_name=display_name,
        group=group,
        method="p2_p98_minmax",
        direction=direction,
        raw_unit=raw_unit,
        raw_path=raw_path,
        score_path=score_path,
        lower_bound=low,
        upper_bound=high,
        formula=formula,
        rule=f"p2={format_number(low)}; p98={format_number(high)}",
    )


def format_number(value: float) -> str:
    magnitude = abs(value)
    if magnitude >= 1000:
        return f"{value:,.0f}"
    if magnitude >= 10:
        return f"{value:.2f}"
    if magnitude >= 1:
        return f"{value:.3f}"
    return f"{value:.4f}"


def public_raw_unit(spec: ComponentSpec) -> str:
    return PUBLIC_RAW_UNITS.get(spec.component_id, spec.raw_unit)


def public_formula(spec: ComponentSpec) -> str:
    if spec.component_id == "built_agriculture_road_expansion":
        return "0.50 × maximum(built-up, agriculture) + 0.50 × road proximity"
    if spec.component_id == "land_cover_pressure":
        return "score = lookup table(class)"
    if spec.component_id in {"bare_soil_vegetation_loss", "critical_infrastructure"}:
        selected_class = "60" if spec.component_id == "bare_soil_vegetation_loss" else "50"
        return f"score = 1[class = {selected_class}]"
    if spec.component_id == "protected_ecological_context":
        return "score = 0 only when absence is verified in the queried source"
    if spec.method == "fixed_distance_decay":
        return "clip(1 - distance / D, 0, 1)"
    return (
        spec.formula.replace("p2", "p_2")
        .replace("p98", "p_98")
    )


def public_rule(spec: ComponentSpec) -> str:
    if spec.method == "p2_p98_minmax":
        return (
            f"2nd percentile={format_number(spec.lower_bound or 0)}; "
            f"98th percentile={format_number(spec.upper_bound or 0)}"
        )
    if spec.component_id == "land_cover_pressure":
        return "Land-cover classes are assigned values using an explicit lookup table."
    if spec.component_id == "bare_soil_vegetation_loss":
        return "Bare or sparse vegetation class 60 → 1; other mapped classes → 0."
    if spec.component_id == "built_agriculture_road_expansion":
        return "0.50 × maximum(built-up, agriculture) + 0.50 × road proximity."
    if spec.component_id == "water_wetland_context":
        return "Linear proximity scaling from mapped water and wetland classes."
    if spec.component_id == "protected_ecological_context":
        return (
            "No protected-area features found in the queried national dataset within 25 km; "
            "score set to 0 for this example."
        )
    if spec.component_id == "road_access_proximity":
        return "Linear proximity scaling; score reaches 0 at 20 km."
    if spec.component_id == "critical_infrastructure":
        return "Built-up class 50 → 1; other mapped classes → 0."
    return spec.rule.replace("LUT", "lookup table")


def build_specs() -> list[ComponentSpec]:
    rainfall_manifest = json.loads(
        (BASE / "rainfall_criterion/rainfall_manifest.json").read_text(encoding="utf-8")
    )
    rainfall_bounds = (
        float(rainfall_manifest["normalization"]["low_mm"]),
        float(rainfall_manifest["normalization"]["high_mm"]),
    )
    exposure_manifest = json.loads(
        (BASE / "exposure_indicators/exposure_manifest.json").read_text(encoding="utf-8")
    )
    population_bounds = (
        float(exposure_manifest["population_normalization"]["low"]),
        float(exposure_manifest["population_normalization"]["high"]),
    )

    continuous = [
        continuous_spec(
            "rainfall_climatology",
            "Rainfall climatology",
            "Flood",
            BASE / "rainfall_criterion/rainfall_dec_mar_2014_2024_mean_mm_30m_resample.tif",
            CRITERIA / "rainfall_climatology_0_1.tif",
            "mm per Dec-Mar season",
            bounds=rainfall_bounds,
        ),
        continuous_spec(
            "twi_wetness",
            "Topographic Wetness Index (TWI)",
            "Flood",
            HYDRO / f"secocha_twi.tif",
            CRITERIA / "twi_wetness_0_1.tif",
            "dimensionless",
        ),
        continuous_spec(
            "drainage_proximity",
            "Drainage proximity",
            "Flood",
            HYDRO / f"secocha_distance_to_drainage_t500.tif",
            CRITERIA / "drainage_proximity_0_1.tif",
            "m",
            invert=True,
        ),
        continuous_spec(
            "soil_k_factor",
            "Soil erodibility (K)",
            "Sediment",
            BASE / "soil_k_factor/soil_k_factor_aligned_250m_effective_30m.tif",
            CRITERIA / "soil_k_factor_0_1.tif",
            "t ha h / (ha MJ mm) [SI]",
        ),
        continuous_spec(
            "ls_factor",
            "Slope–length factor (LS)",
            "Sediment",
            BASE / "ls_factor" / f"secocha_ls_factor_capped50_hillslope_30m.tif",
            CRITERIA / "ls_factor_0_1.tif",
            "dimensionless",
        ),
        continuous_spec(
            "c_factor",
            "Cover-management factor (C)",
            "Sediment",
            BASE / "worldcover_factors/c_factor_30m.tif",
            CRITERIA / "c_factor_0_1.tif",
            "dimensionless",
        ),
    ]

    population = ComponentSpec(
        component_id="settlement_population",
        display_name="Settlement population",
        group="Socioeconomic",
        method="p2_p98_minmax",
        direction="higher_raw_higher_score",
        raw_unit="people per 30 m cell",
        score_path=BASE / "exposure_indicators/population_exposure_0_1.tif",
        raw_path=ROOT / "data/raw/worldpop_peru_population/per_ppp_2020_UNadj.tif",
        lower_bound=population_bounds[0],
        upper_bound=population_bounds[1],
        formula="clip((x - p2) / (p98 - p2), 0, 1)",
        rule=(
            f"p2={format_number(population_bounds[0])}; "
            f"p98={format_number(population_bounds[1])}"
        ),
    )

    other = [
        ComponentSpec(
            "land_cover_pressure",
            "Land-cover runoff pressure",
            "Land use",
            "class_value_lookup",
            "class_specific_rule",
            "WorldCover class code (categorical)",
            CRITERIA / "land_cover_pressure_0_1.tif",
            "WorldCover class -> explicit runoff-pressure value",
            raw_path=ROOT / "data/raw/esa_worldcover_2021_peru_former/ESA_WorldCover_10m_2021_v200_S18W075_Map.tif",
            formula="score = LUT(WorldCover class)",
        ),
        ComponentSpec(
            "bare_soil_vegetation_loss",
            "Bare-soil / vegetation loss",
            "Land use",
            "binary_class_lookup",
            "class_specific_rule",
            "WorldCover class code (categorical)",
            CRITERIA / "bare_soil_vegetation_loss_0_1.tif",
            "bare/sparse class 60 -> 1; other mapped classes -> 0",
            raw_path=ROOT / "data/raw/esa_worldcover_2021_peru_former/ESA_WorldCover_10m_2021_v200_S18W075_Map.tif",
            formula="score = 1[class = 60]",
        ),
        ComponentSpec(
            "built_agriculture_road_expansion",
            "Built-up, agriculture and road pressure",
            "Land use",
            "weighted_composite",
            "higher_component_higher_score",
            "unitless 0-1 component scores",
            CRITERIA / "built_agriculture_road_expansion_0_1.tif",
            "0.50 x max(built-up, agriculture) + 0.50 x road proximity",
            formula="0.50 max(B, A) + 0.50 R",
        ),
        ComponentSpec(
            "water_wetland_context",
            "Water and wetland proximity",
            "Environmental",
            "fixed_distance_decay",
            "lower_raw_higher_score",
            "m from WorldCover class 80/90",
            BASE / "environmental_sensitivity/water_wetland_context_0_1.tif",
            "linear proximity decay from mapped water/wetland cells",
            raw_path=ROOT / "data/raw/esa_worldcover_2021_peru_former/ESA_WorldCover_10m_2021_v200_S18W075_Map.tif",
            lower_bound=0.0,
            upper_bound=10_000.0,
            formula="clip(1 - distance / 10000, 0, 1)",
        ),
        ComponentSpec(
            "protected_ecological_context",
            "Protected-area context",
            "Environmental",
            "confirmed_zero_context",
            "confirmed_absence_in_queried_layer",
            "feature presence (0/1)",
            BASE / "environmental_sensitivity/protected_area_context_0_1.tif",
            "0 confirmed features in national layer 0 within 25 km; score 0",
            raw_path=ROOT / "data/raw/sernanp_protected_areas_secocha/secocha_sernanp_layer0_query.geojson",
            lower_bound=0.0,
            upper_bound=0.0,
            formula="score = 0 (confirmed queried-layer absence)",
        ),
        ComponentSpec(
            "road_access_proximity",
            "Road proximity",
            "Socioeconomic",
            "fixed_distance_decay",
            "lower_raw_higher_score",
            "m from OSM road",
            CRITERIA / "road_access_proximity_0_1.tif",
            "linear proximity decay; zero at 20 km",
            raw_path=ROOT / "data/raw/geofabrik_osm_peru_former/extracted/gis_osm_roads_free_1.shp",
            lower_bound=0.0,
            upper_bound=20_000.0,
            formula="clip(1 - distance / 20000, 0, 1)",
        ),
        ComponentSpec(
            "critical_infrastructure",
            "Built-up infrastructure",
            "Socioeconomic",
            "binary_class_lookup",
            "class_specific_rule",
            "WorldCover class code (categorical)",
            BASE / "exposure_indicators/built_infrastructure_exposure_0_1.tif",
            "WorldCover built-up class 50 -> 1; other mapped classes -> 0",
            raw_path=ROOT / "data/raw/esa_worldcover_2021_peru_former/ESA_WorldCover_10m_2021_v200_S18W075_Map.tif",
            formula="score = 1[class = 50]",
        ),
    ]
    return [*continuous, *other[:3], *other[3:5], other[5], population, other[6]]


def row_from_spec(spec: ComponentSpec, boundary_cells: int) -> dict[str, str | float]:
    score = read_array(spec.score_path)
    valid = score[np.isfinite(score)]
    if not valid.size:
        raise ValueError(f"no valid standardized cells: {spec.score_path}")
    return {
        "component_id": spec.component_id,
        "display_name": spec.display_name,
        "criterion_group": spec.group,
        "method": PUBLIC_METHOD_LABELS[spec.method],
        "direction": PUBLIC_DIRECTION_LABELS[spec.direction],
        "raw_unit": public_raw_unit(spec),
        "lower_bound": "" if spec.lower_bound is None else f"{spec.lower_bound:.10g}",
        "upper_bound": "" if spec.upper_bound is None else f"{spec.upper_bound:.10g}",
        "formula": public_formula(spec),
        "rule_detail": public_rule(spec),
        "score_min": f"{float(valid.min()):.10g}",
        "score_max": f"{float(valid.max()):.10g}",
        "score_mean": f"{float(valid.mean()):.10g}",
        "valid_fraction": f"{valid.size / boundary_cells:.10g}",
        "claim_level": PUBLIC_STATUS,
        "raw_source_path": PUBLIC_SOURCE_LABELS.get(spec.component_id, "Source dataset description"),
        "score_path": PUBLIC_OUTPUT_LABEL,
    }


def source_registry_rows(rows: list[dict[str, str | float]]) -> list[dict[str, str]]:
    """Group indicators into a compact public source register."""
    used_for: dict[str, list[str]] = {key: [] for key in PUBLIC_SOURCE_REGISTRY}
    for row in rows:
        component_id = str(row["component_id"])
        source_key = PUBLIC_SOURCE_KEYS[component_id]
        used_for[source_key].append(str(row["display_name"]))
    return [
        {
            **record,
            "used_for": "; ".join(used_for[source_key]),
            "claim_level": PUBLIC_STATUS,
        }
        for source_key, record in PUBLIC_SOURCE_REGISTRY.items()
    ]


def write_component_csv(path: Path, rows: list[dict[str, str | float]]) -> None:
    fields = [label for label, _ in EXPLANATION_COLUMNS]
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(
            {
                label: row[field]
                for label, field in EXPLANATION_COLUMNS
            }
            for row in rows
        )


def add_method_cards(ax: plt.Axes) -> tuple[list[plt.Text], list[FancyBboxPatch]]:
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    cards = [
        (
            0.01,
            "1  Continuous values",
            "2nd–98th percentile clipping + min–max scaling\nscore = clip((x-p_2)/(p_98-p_2), 0, 1)\nReverse direction when lower values indicate higher priority",
            "#dbeafe",
        ),
        (
            0.26,
            "2  Distance-decay scaling",
            "score = clip(1 - distance / D, 0, 1)\nThe maximum distance D is recorded for each feature type",
            "#dcfce7",
        ),
        (
            0.51,
            "3  Categorical rules",
            "Class code → explicit lookup table or 0/1 rule\nThe coding system and mapping are recorded",
            "#fef3c7",
        ),
        (
            0.76,
            "4  Combination + missing-data rule",
            "Combine only declared 0–1 indicators\nOutside-study-area and missing inputs remain NoData",
            "#f3e8ff",
        ),
    ]
    text_artists: list[plt.Text] = []
    card_patches: list[FancyBboxPatch] = []
    for x, title, body, color in cards:
        patch = FancyBboxPatch(
            (x, 0.14),
            0.225,
            0.72,
            boxstyle="round,pad=0.012,rounding_size=0.018",
            linewidth=1.2,
            edgecolor="#475569",
            facecolor=color,
            transform=ax.transAxes,
        )
        ax.add_patch(patch)
        card_patches.append(patch)
        text_artists.append(
            ax.text(
                x + 0.014,
                0.70,
                title,
                transform=ax.transAxes,
                fontsize=10.5,
                fontweight="bold",
            )
        )
        text_artists.append(ax.text(
            x + 0.014,
            0.56,
            body,
            transform=ax.transAxes,
            fontsize=8.5,
            va="top",
            linespacing=1.35,
            color="#263238",
        ))
    return text_artists, card_patches


def add_transfer_plot(ax: plt.Axes, spec: ComponentSpec) -> plt.Text:
    if spec.raw_path is None or spec.lower_bound is None or spec.upper_bound is None:
        raise ValueError(f"transfer plot requires raw path and bounds: {spec.component_id}")
    raw = read_array(spec.raw_path)
    score = read_array(spec.score_path)
    valid = np.isfinite(raw) & np.isfinite(score)
    x_values = raw[valid]
    y_values = score[valid]
    order = np.argsort(x_values)
    if order.size > 700:
        order = order[np.linspace(0, order.size - 1, 700).astype(int)]
    ax.scatter(x_values[order], y_values[order], s=6, alpha=0.18, color="#334155", linewidths=0)
    x_line = np.linspace(float(np.nanmin(x_values)), float(np.nanmax(x_values)), 300)
    y_line = np.clip(
        (x_line - spec.lower_bound) / (spec.upper_bound - spec.lower_bound), 0, 1
    )
    if spec.direction == "lower_raw_higher_score":
        y_line = 1 - y_line
    ax.plot(x_line, y_line, color="#dc2626", linewidth=2.2)
    ax.axvline(spec.lower_bound, color="#2563eb", linestyle="--", linewidth=1)
    ax.axvline(spec.upper_bound, color="#2563eb", linestyle="--", linewidth=1)
    ax.text(spec.lower_bound, 1.035, "p2", ha="center", va="bottom", fontsize=7, color="#2563eb")
    ax.text(spec.upper_bound, 1.035, "p98", ha="center", va="bottom", fontsize=7, color="#2563eb")
    ax.set_ylim(-0.04, 1.09)
    title = ax.set_title(spec.display_name, fontsize=10.5, fontweight="bold", pad=7)
    ax.set_xlabel(public_raw_unit(spec), fontsize=8)
    ax.set_ylabel("Standardized score (0–1)", fontsize=8)
    ax.tick_params(labelsize=7)
    ax.grid(axis="y", color="#e2e8f0", linewidth=0.7)
    ax.spines[["top", "right"]].set_visible(False)
    return title


def cross_row_layout_metrics(
    fig: plt.Figure,
    upper_texts: list[plt.Text],
    card_patches: list[FancyBboxPatch],
    lower_texts: list[plt.Text],
    title: plt.Text | None = None,
    subtitle: plt.Text | None = None,
    footer: plt.Text | None = None,
    xlabels: list[plt.Text] | None = None,
) -> dict[str, int | float]:
    """Measure unintended collisions and clearance across adjacent rows."""
    fig.canvas.draw()
    renderer = fig.canvas.get_renderer()
    upper_boxes = [artist.get_window_extent(renderer=renderer) for artist in upper_texts]
    card_boxes = [artist.get_window_extent(renderer=renderer) for artist in card_patches]
    lower_boxes = [artist.get_window_extent(renderer=renderer) for artist in lower_texts]
    text_overlaps = sum(
        upper.overlaps(lower) for upper in upper_boxes for lower in lower_boxes
    )
    card_overlaps = sum(
        card.overlaps(lower) for card in card_boxes for lower in lower_boxes
    )
    min_gap = min(card.y0 - lower.y1 for card in card_boxes for lower in lower_boxes)
    card_overflow_count = 0
    for index, card_box in enumerate(card_boxes):
        for text_box in upper_boxes[index * 2 : index * 2 + 2]:
            if (
                text_box.x0 < card_box.x0 - 1
                or text_box.x1 > card_box.x1 + 1
                or text_box.y0 < card_box.y0 - 1
                or text_box.y1 > card_box.y1 + 1
            ):
                card_overflow_count += 1
    overlap_count = int(text_overlaps + card_overlaps)
    checked_text_count = len(upper_texts) + len(lower_texts)
    metrics: dict[str, int | float] = {
        "layout_overlap_count": overlap_count,
        "layout_card_overflow_count": card_overflow_count,
        "layout_checked_text_count": checked_text_count,
        "layout_min_vertical_gap_px": round(float(min_gap), 2),
    }
    if title is not None and subtitle is not None:
        title_box = title.get_window_extent(renderer=renderer)
        subtitle_box = subtitle.get_window_extent(renderer=renderer)
        metrics["layout_title_subtitle_gap_px"] = round(float(title_box.y0 - subtitle_box.y1), 2)
        overlap_count += int(title_box.overlaps(subtitle_box))
        checked_text_count += 2
    if footer is not None and xlabels:
        footer_box = footer.get_window_extent(renderer=renderer)
        xlabel_boxes = [label.get_window_extent(renderer=renderer) for label in xlabels]
        metrics["layout_footer_gap_px"] = round(
            float(min(label.y0 - footer_box.y1 for label in xlabel_boxes)), 2
        )
        overlap_count += sum(footer_box.overlaps(label_box) for label_box in xlabel_boxes)
        checked_text_count += 1 + len(xlabel_boxes)
    metrics["layout_overlap_count"] = int(overlap_count + card_overflow_count)
    metrics["layout_checked_text_count"] = int(checked_text_count)
    return metrics


def compact_method(spec: ComponentSpec) -> str:
    labels = {
        "p2_p98_minmax": "percentile min–max",
        "fixed_distance_decay": "distance decay",
        "class_value_lookup": "categorical lookup table",
        "binary_class_lookup": "binary class rule",
        "weighted_composite": "weighted combination",
        "confirmed_zero_context": "verified absence",
    }
    return labels[spec.method]


def compact_direction(spec: ComponentSpec) -> str:
    labels = {
        "higher_raw_higher_score": "higher measurement → higher score",
        "lower_raw_higher_score": "lower measurement → higher score",
        "class_specific_rule": "class-based interpretation",
        "higher_component_higher_score": "higher indicators → higher combined score",
        "confirmed_absence_in_queried_layer": "verified absence → 0",
    }
    return labels[spec.direction]


def compact_raw_unit(spec: ComponentSpec) -> str:
    return public_raw_unit(spec)


def render_process_figure(
    path_png: Path, path_svg: Path, specs: list[ComponentSpec]
) -> dict[str, int]:
    fig = plt.figure(figsize=(19, 7.5), facecolor="white")
    grid = fig.add_gridspec(2, 1, height_ratios=[1.0, 2.4], hspace=0.58)
    cards_ax = fig.add_subplot(grid[0])
    card_texts, card_patches = add_method_cards(cards_ax)

    examples = fig.add_subfigure(grid[1]).subplots(1, 4)
    by_id = {spec.component_id: spec for spec in specs}
    example_titles: list[plt.Text] = []
    example_xlabels: list[plt.Text] = []
    for ax, component_id in zip(
        examples,
        ["rainfall_climatology", "twi_wetness", "drainage_proximity", "soil_k_factor"],
    ):
        example_titles.append(add_transfer_plot(ax, by_id[component_id]))
        example_xlabels.append(ax.xaxis.label)

    title = fig.suptitle(
        "Standardization process: original measurements to comparable 0–1 scores",
        x=0.035,
        y=0.985,
        ha="left",
        fontsize=19,
        fontweight="bold",
    )
    subtitle = fig.text(
        0.035,
        0.910,
        "The figure shows the transformation families and four example relationships; the companion Excel file contains the full indicator table.",
        fontsize=10.5,
        color="#334155",
    )
    footer = fig.text(
        0.035,
        0.025,
        "Preliminary example | 30-m reference grid | Missing-data cells remain NoData and are never assigned a score of zero.",
        fontsize=8.5,
        color="#475569",
    )
    fig.subplots_adjust(left=0.035, right=0.985, top=0.84, bottom=0.19)
    layout = cross_row_layout_metrics(
        fig,
        card_texts,
        card_patches,
        example_titles,
        title=title,
        subtitle=subtitle,
        footer=footer,
        xlabels=example_xlabels,
    )
    fig.savefig(path_png, dpi=200, facecolor="white")
    fig.savefig(path_svg, facecolor="white")
    plt.close(fig)
    return layout


def write_excel(path: Path, rows: list[dict], study_area: str = STUDY_AREA) -> None:
    """Write a compact detail table with separate explanations and rule sheets."""
    if Workbook is None:
        raise SystemExit("openpyxl is needed for the xlsx audit table: .venv/bin/python -m pip install openpyxl")
    workbook = Workbook()
    readme = workbook.active
    readme.title = "Read me"
    sheet = workbook.create_sheet("Standardization detail")
    explanations = workbook.create_sheet("Explanations")
    sources = workbook.create_sheet("Data sources")
    additional = workbook.create_sheet("Additional materials")
    rules = workbook.create_sheet("Rule catalog")

    font_name = "Arial"
    title_fill = PatternFill("solid", fgColor="DDEBF7")
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    title_font = Font(name=font_name, size=16, bold=True, color="1F4E78")
    header_font = Font(name=font_name, bold=True, color="1F4E78")
    note_font = Font(name=font_name, color="44546A")
    body_font = Font(name=font_name, size=10, color="000000")
    hyperlink_font = Font(name=font_name, size=10, color="0563C1", underline="single")
    header_row = 4

    def configure_sheet(sheet_obj, last_col: str, end_row: int, widths: list[int], *, paper: str = "A3") -> None:
        sheet_obj.sheet_view.showGridLines = True
        sheet_obj.freeze_panes = "A5"
        sheet_obj.print_area = f"A1:{last_col}{end_row}"
        sheet_obj.print_title_rows = "1:4"
        sheet_obj.page_setup.orientation = "landscape"
        sheet_obj.page_setup.paperSize = sheet_obj.PAPERSIZE_A3 if paper == "A3" else sheet_obj.PAPERSIZE_A4
        sheet_obj.page_setup.fitToWidth = 1
        sheet_obj.page_setup.fitToHeight = 0 if paper == "A3" else 1
        sheet_obj.sheet_properties.pageSetUpPr.fitToPage = True
        sheet_obj.page_margins.left = 0.25
        sheet_obj.page_margins.right = 0.25
        sheet_obj.page_margins.top = 0.35
        sheet_obj.page_margins.bottom = 0.35
        for index, width in enumerate(widths, start=1):
            sheet_obj.column_dimensions[get_column_letter(index)].width = width

    def style_title_and_notes(sheet_obj, last_col: str, title: str, note1: str, note2: str) -> None:
        sheet_obj.merge_cells(f"A1:{last_col}1")
        sheet_obj["A1"] = title
        sheet_obj["A1"].font = title_font
        sheet_obj["A1"].fill = title_fill
        sheet_obj.merge_cells(f"A2:{last_col}2")
        sheet_obj["A2"] = note1
        sheet_obj["A2"].font = Font(name=font_name, italic=True, color="44546A")
        sheet_obj.merge_cells(f"A3:{last_col}3")
        sheet_obj["A3"] = note2
        sheet_obj["A3"].font = note_font
        sheet_obj.row_dimensions[1].height = 25
        sheet_obj.row_dimensions[2].height = 22
        sheet_obj.row_dimensions[3].height = 30

    def style_headers(sheet_obj, columns: list[tuple[str, str]] | list[str]) -> None:
        for column_index, item in enumerate(columns, start=1):
            label = item[0] if isinstance(item, tuple) else item
            cell = sheet_obj.cell(row=header_row, column=column_index, value=label)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet_obj.row_dimensions[header_row].height = 32

    def add_light_table(sheet_obj, display_name: str, last_col: str, end_row: int) -> None:
        table = Table(displayName=display_name, ref=f"A{header_row}:{last_col}{end_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleLight1",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=False,
            showColumnStripes=False,
        )
        sheet_obj.add_table(table)

    # Compact overview: the original populated table, without long explanations.
    summary_last_col = get_column_letter(len(SUMMARY_COLUMNS))
    style_title_and_notes(
        sheet,
        summary_last_col,
        "Standardization detail",
        f"Study area: {study_area} | Evidence status: {PUBLIC_STATUS}",
        "Region-independent overview of the populated example. Formulas, source descriptions, and standardized output details are in Explanations.",
    )
    style_headers(sheet, SUMMARY_COLUMNS)
    summary_numeric_fields = {
        "lower_bound", "upper_bound", "score_min", "score_max", "score_mean", "valid_fraction"
    }
    for row_index, row in enumerate(rows, start=header_row + 1):
        for column_index, (_, field) in enumerate(SUMMARY_COLUMNS, start=1):
            value = row[field]
            if field in summary_numeric_fields:
                value = None if value == "" else float(value)
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if field in {"lower_bound", "upper_bound"}:
                cell.number_format = "0.000000"
            elif field in {"score_min", "score_max", "score_mean"}:
                cell.number_format = "0.000"
            elif field == "valid_fraction":
                cell.number_format = "0.0%"
        sheet.row_dimensions[row_index].height = 32
    summary_end_row = header_row + len(rows)
    add_light_table(sheet, "StandardizationDetail", summary_last_col, summary_end_row)
    configure_sheet(sheet, summary_last_col, summary_end_row, [34, 18, 32, 28, 36, 30, 14, 14, 16, 16, 16, 16, 20])

    # Long text is deliberately separated so the overview remains readable.
    explanation_last_col = get_column_letter(len(EXPLANATION_SHEET_COLUMNS))
    style_title_and_notes(
        explanations,
        explanation_last_col,
        "Explanations",
        "Long text kept separate from the compact detail table for easier reading.",
        "This sheet records the formula, rule description, source dataset, and standardized output for each indicator.",
    )
    style_headers(explanations, EXPLANATION_SHEET_COLUMNS)
    for row_index, row in enumerate(rows, start=header_row + 1):
        for column_index, (_, field) in enumerate(EXPLANATION_SHEET_COLUMNS, start=1):
            cell = explanations.cell(row=row_index, column=column_index, value=row[field])
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        explanations.row_dimensions[row_index].height = 54
    explanation_end_row = header_row + len(rows)
    add_light_table(explanations, "StandardizationExplanations", explanation_last_col, explanation_end_row)
    configure_sheet(explanations, explanation_last_col, explanation_end_row, [24, 34, 44, 64, 48, 40])

    # Public source register: one traceable source description for each indicator.
    source_last_col = get_column_letter(len(SOURCE_COLUMNS))
    style_title_and_notes(
        sources,
        source_last_col,
        "Data sources",
        "Public source descriptions used by the indicators in this example.",
        "Each row groups indicators that use the same source; formulas and rules remain in Explanations.",
    )
    style_headers(sources, SOURCE_COLUMNS)
    source_rows = source_registry_rows(rows)
    for row_index, row in enumerate(source_rows, start=header_row + 1):
        for column_index, (_, field) in enumerate(SOURCE_COLUMNS, start=1):
            cell = sources.cell(row=row_index, column=column_index, value=row[field])
            cell.font = hyperlink_font if field in {"source_portal", "download_url"} else body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if field in {"source_portal", "download_url"}:
                urls = str(row[field]).splitlines()
                if urls and urls[0].startswith("https://"):
                    cell.hyperlink = urls[0]
        sources.row_dimensions[row_index].height = 64
    source_end_row = header_row + len(source_rows)
    add_light_table(sources, "StandardizationDataSources", source_last_col, source_end_row)
    configure_sheet(sources, source_last_col, source_end_row, [14, 42, 78, 62, 88, 86, 20])

    # Supplementary materials are kept separate from the raw indicator sources.
    additional_last_col = get_column_letter(len(ADDITIONAL_COLUMNS))
    style_title_and_notes(
        additional,
        additional_last_col,
        "Additional materials",
        "Supporting materials used by or produced by the current model chain.",
        "This sheet separates spatial context, model parameters, and derived outputs from downloadable indicator sources.",
    )
    style_headers(additional, ADDITIONAL_COLUMNS)
    for row_index, row in enumerate(ADDITIONAL_ROWS, start=header_row + 1):
        for column_index, (_, field) in enumerate(ADDITIONAL_COLUMNS, start=1):
            cell = additional.cell(row=row_index, column=column_index, value=row[field])
            cell.font = hyperlink_font if field == "source_url" and str(row[field]).startswith("https://") else body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if field == "source_url" and str(row[field]).startswith("https://"):
                cell.hyperlink = str(row[field])
        additional.row_dimensions[row_index].height = 58
    additional_end_row = header_row + len(ADDITIONAL_ROWS)
    add_light_table(additional, "StandardizationAdditionalMaterials", additional_last_col, additional_end_row)
    configure_sheet(additional, additional_last_col, additional_end_row, [14, 34, 32, 54, 78, 40, 70])

    # Generic method definitions are kept as a reference sheet, not mixed into the detail table.
    rules.sheet_view.showGridLines = True
    rules.freeze_panes = "A5"
    rule_last_col = get_column_letter(len(RULE_CATALOG_COLUMNS))
    style_title_and_notes(
        rules,
        rule_last_col,
        "Rule catalog",
        "Generic transformation families used to map source values to comparable 0–1 scores.",
        "Each selected rule still requires area-specific parameters, source evidence, and a declared NoData policy.",
    )
    style_headers(rules, RULE_CATALOG_COLUMNS)
    for row_index, rule_row in enumerate(RULE_CATALOG_ROWS, start=header_row + 1):
        for column_index, value in enumerate(rule_row, start=1):
            cell = rules.cell(row=row_index, column=column_index, value=value)
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        rules.row_dimensions[row_index].height = 60
    rule_end_row = header_row + len(RULE_CATALOG_ROWS)
    add_light_table(rules, "StandardizationRuleCatalog", rule_last_col, rule_end_row)
    configure_sheet(rules, rule_last_col, rule_end_row, [24, 28, 38, 38, 48, 32, 32, 42, 40])

    # Reading guide.
    # Keep the reading guide spreadsheet-like: visible gridlines make the two-column
    # label/value relationship easier to follow in Excel.
    readme.sheet_view.showGridLines = True
    readme.merge_cells("A1:B1")
    readme["A1"] = "How to read this workbook"
    readme["A1"].font = Font(name=font_name, size=15, bold=True, color="1F4E78")
    readme["A1"].fill = title_fill
    notes = [
        ("Purpose", "Use Standardization detail for a concise overview; use Explanations for the long text behind each row."),
        ("Standardization detail", f"Current {study_area} example with original units, bounds, methods, and standardized summary values."),
        ("Explanations", "Formula, rule description, source dataset, and standardized output details for each indicator."),
        ("Data sources", "Source dataset, portal, download URL, indicators that use it, and a public processing note."),
        ("Additional materials", "Boundary/context materials, model parameters, and derived outputs that are not raw indicator sources."),
        ("Rule catalog", "Generic standardization methods and the evidence that should accompany each method."),
        ("Study area", study_area),
        ("Study-area definition", "Reference extent used for this example; it is not a hydrological catchment or validated hazard footprint."),
        ("Original unit", "Unit or source coding before conversion to the comparable 0–1 score."),
        ("Missing-data handling", "Cells outside the study area or missing in the source remain NoData; no zero imputation."),
        ("Evidence status", PUBLIC_STATUS + "; not a validated flood-risk result."),
        ("Companion process figure", "Use the process figure generated for the selected study area."),
    ]
    for row_index, (label, value) in enumerate(notes, start=3):
        label_cell = readme.cell(row=row_index, column=1, value=label)
        label_cell.font = Font(name=font_name, bold=True, color="1F4E78")
        label_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        value_cell = readme.cell(row=row_index, column=2, value=value)
        value_cell.font = body_font
        value_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        readme.row_dimensions[row_index].height = 30
    readme.column_dimensions["A"].width = 44
    readme.column_dimensions["B"].width = 90
    readme.print_area = f"A1:B{2 + len(notes)}"
    readme.page_setup.orientation = "landscape"
    readme.page_setup.paperSize = readme.PAPERSIZE_A4
    readme.page_setup.fitToWidth = 1
    readme.page_setup.fitToHeight = 1
    readme.sheet_properties.pageSetUpPr.fitToPage = True
    readme.page_margins.left = 0.35
    readme.page_margins.right = 0.35
    readme.page_margins.top = 0.45
    readme.page_margins.bottom = 0.45

    workbook.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    dem = read_array(BASE / f"terrain_preprocessing/secocha_dem_utm18s_30m.tif")
    boundary_cells = int(np.isfinite(dem).sum())
    if boundary_cells == 0:
        raise ValueError("active Secocha boundary has no valid reference cells")

    specs = build_specs()
    rows = [row_from_spec(spec, boundary_cells) for spec in specs]
    source_rows = source_registry_rows(rows)
    csv_path = output_dir / "standardization_detail.csv"
    process_png = output_dir / "secocha_standardization_process.png"
    process_svg = output_dir / "secocha_standardization_process.svg"
    xlsx_path = output_dir / "IFPM_Standardization_Audit.xlsx"
    manifest_path = output_dir / "standardization_detail_manifest.json"
    write_component_csv(csv_path, rows)
    layout = render_process_figure(process_png, process_svg, specs)
    write_excel(xlsx_path, rows, STUDY_AREA)

    method_counts: dict[str, int] = {}
    for spec in specs:
        method_label = PUBLIC_METHOD_LABELS[spec.method]
        method_counts[method_label] = method_counts.get(method_label, 0) + 1
    manifest = {
        "figure_id": "standardization_detail",
        "public_safe": True,
        "study_area": STUDY_AREA,
        "status": PUBLIC_STATUS,
        "indicator_count": len(specs),
        "method_counts": method_counts,
        "separated_outputs": True,
        "generic_table_artifact": True,
        "table_artifact_name": xlsx_path.name,
        "table_format": "xlsx",
        "example_worksheet": "Standardization detail",
        "explanation_worksheet": "Explanations",
        "explanation_row_count": len(rows),
        "source_worksheet": "Data sources",
        "source_row_count": len(source_rows),
        "source_download_url_column": "Download URL",
        "additional_worksheet": "Additional materials",
        "additional_row_count": len(ADDITIONAL_ROWS),
        "rule_catalog_worksheet": "Rule catalog",
        "rule_catalog_count": len(RULE_CATALOG_ROWS),
        "takeaway": "different indicators require different standardization methods; each method and interpretation is recorded",
        "nodata_policy": "missing and outside-study-area cells remain NoData; no zero imputation",
        "visible_table_columns": [label for label, _ in SUMMARY_COLUMNS],
        "figures": {
            "process": {
                "png": process_png.name,
                "svg": process_svg.name,
                **layout,
            },
            "table": {
                "xlsx": xlsx_path.name,
                "worksheet": "Standardization detail",
                "row_count": len(rows),
                "explanation_worksheet": "Explanations",
                "explanation_row_count": len(rows),
                "source_worksheet": "Data sources",
                "source_row_count": len(source_rows),
                "additional_worksheet": "Additional materials",
                "additional_row_count": len(ADDITIONAL_ROWS),
            },
        },
        **layout,
        "outputs": [path.name for path in (process_png, process_svg, xlsx_path, csv_path)],
    }
    manifest_path.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(manifest, indent=2))


if __name__ == "__main__":
    main()
