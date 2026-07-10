#!/usr/bin/env python3
"""Extract Arequipa district poverty records from the INEI 2018 poverty workbook."""

from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
RAW_DIR = ROOT / "data/raw/inei_peru_poverty_2018"
OUT_DIR = ROOT / "outputs/inei_peru_poverty_2018"
WORKBOOK = RAW_DIR / "anexo_estadistico_mapa_pobreza_2018.xlsx"
AREQUIPA_CSV = RAW_DIR / "arequipa_district_poverty_2018.csv"

INPUTS = [ROOT / "data/raw/inei_peru_poverty_2018/anexo_estadistico_mapa_pobreza_2018.xlsx"]
OUTPUTS = [
    ROOT / "data/raw/inei_peru_poverty_2018/arequipa_district_poverty_2018.csv",
    ROOT / "outputs/inei_peru_poverty_2018/inei_arequipa_poverty_manifest.json",
]


def rel(path: Path) -> str:
    return str(path.resolve().relative_to(ROOT))


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def main() -> None:
    if not WORKBOOK.exists():
        raise FileNotFoundError(f"missing INEI workbook: {rel(WORKBOOK)}")
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
    ws = wb["A3.4"]  # AREQUIPA department district table

    rows = []
    for row in ws.iter_rows(min_row=8, values_only=True):
        ubigeo = row[0]
        if ubigeo is None:
            continue
        ubigeo = str(ubigeo).strip()
        if not ubigeo or not ubigeo.isdigit():
            continue
        rows.append(
            {
                "ubigeo": ubigeo.zfill(6),
                "sufijo_distrito": str(row[1]).strip() if row[1] is not None else "",
                "provincia": str(row[2]).strip() if row[2] is not None else "",
                "distrito": str(row[3]).strip() if row[3] is not None else "",
                "poblacion_proyectada_2020": row[4],
                "pobreza_total_ci95_inferior": row[5],
                "pobreza_total_ci95_superior": row[6],
                "grupo_robusto": row[7],
                "ubicacion_pobreza_monetaria_total": row[8],
            }
        )

    with AREQUIPA_CSV.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)

    lower_vals = [float(r["pobreza_total_ci95_inferior"]) for r in rows if r["pobreza_total_ci95_inferior"] is not None]
    upper_vals = [float(r["pobreza_total_ci95_superior"]) for r in rows if r["pobreza_total_ci95_superior"] is not None]
    summary = {
        "dataset_id": "inei_peru_poverty_2018",
        "source_workbook": rel(WORKBOOK),
        "source_workbook_size_bytes": WORKBOOK.stat().st_size,
        "source_workbook_sha256": sha256_file(WORKBOOK),
        "arequipa_csv": rel(AREQUIPA_CSV),
        "arequipa_csv_size_bytes": AREQUIPA_CSV.stat().st_size,
        "arequipa_csv_sha256": sha256_file(AREQUIPA_CSV),
        "district_records": len(rows),
        "poverty_ci95_lower_min": min(lower_vals) if lower_vals else None,
        "poverty_ci95_lower_mean": sum(lower_vals) / len(lower_vals) if lower_vals else None,
        "poverty_ci95_lower_max": max(lower_vals) if lower_vals else None,
        "poverty_ci95_upper_min": min(upper_vals) if upper_vals else None,
        "poverty_ci95_upper_mean": sum(upper_vals) / len(upper_vals) if upper_vals else None,
        "poverty_ci95_upper_max": max(upper_vals) if upper_vals else None,
        "claim_level": "diagnostic_only",
        "evidence_role": "district_response_capacity_context",
        "note": "District-level monetary poverty context; not pixel-scale adaptive capacity by itself.",
    }
    manifest_path = OUT_DIR / "inei_arequipa_poverty_manifest.json"
    with manifest_path.open("w", encoding="utf-8") as fh:
        json.dump(summary, fh, indent=2, ensure_ascii=False)

    print(f"wrote: {rel(AREQUIPA_CSV)}")
    print(f"wrote: {rel(manifest_path)}")
    print(f"district records: {len(rows)}")


if __name__ == "__main__":
    main()
